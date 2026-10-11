"""Excel export with basic professional styling."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Union

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def create_styled_excel(
    data: Union[pd.DataFrame, Dict[str, pd.DataFrame]],
    output_path: str,
) -> str:
    """Export DataFrame(s) to Excel with styled headers and zebra striping."""
    out = Path(output_path).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)

    sheets: Dict[str, pd.DataFrame]
    if isinstance(data, pd.DataFrame):
        sheets = {"Sheet1": data}
    else:
        sheets = {name: df for name, df in data.items()}

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31] or "Sheet1"
            df.to_excel(writer, sheet_name=safe_name, index=False)
            ws = writer.sheets[safe_name]

            header_fill = PatternFill("solid", fgColor="1F4E79")
            header_font = Font(color="FFFFFF", bold=True)
            zebra_fill = PatternFill("solid", fgColor="F2F2F2")

            for col_idx, _ in enumerate(df.columns, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row_idx in range(2, ws.max_row + 1):
                if row_idx % 2 == 0:
                    for col_idx in range(1, ws.max_column + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = zebra_fill

            for col_idx, column in enumerate(df.columns, start=1):
                values = [str(column)] + [
                    str(v) if v is not None else "" for v in df[column].tolist()
                ]
                width = min(max(len(v) for v in values) + 2, 60)
                ws.column_dimensions[get_column_letter(col_idx)].width = width

    return str(out)
